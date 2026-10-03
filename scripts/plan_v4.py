# -*- coding: utf-8 -*-
# ==== L4D2 skill: 统一路径解析（构建时注入，勿手改）====
import os as _os, json as _json
def _l4d2_home():
    h = _os.environ.get("L4D2_HOME")
    if h:
        return h
    return _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
L4D2_HOME = _l4d2_home()
def _l4d2_cfg():
    try:
        with open(_os.path.join(L4D2_HOME, "config.json"), encoding="utf-8") as f:
            return _json.load(f)
    except Exception:
        return {}
L4D2_CFG = _l4d2_cfg()
L4D2_WORK = _os.path.join(L4D2_HOME, L4D2_CFG.get("scriptsDir", "_work"))
L4D2_XLSX = (_os.environ.get("L4D2_XLSX_OUT")
             or _os.path.join(L4D2_HOME, L4D2_CFG.get("xlsxName", "求生之路2_可MOD替换物品总表.xlsx")))
L4D2_ADDONS = (_os.environ.get("L4D2_ADDONS")
               or L4D2_CFG.get("addonsDir")
               or r"D:\Steam\steamapps\common\Left 4 Dead 2\left4dead2\addons")
# ==== 注入结束 ====
import sys, json, os, re
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

W = L4D2_WORK
XLSX = L4D2_XLSX
OUT = os.path.join(W, 'rename_plan.md')

scan = json.load(open(os.path.join(W, 'addons_scan.json'), encoding='utf-8'))
# ⚠️ 扫描数据过期检测：改名之后如果没重跑「刷新表格.bat」，addons_scan.json 里还是**旧文件名**，
#    本推演会对着旧名再提一次同样的改名 —— 看起来像「幂等失效」，其实只是数据旧了。
_stale = [m['file'] for m in scan
          if m.get('file') and not os.path.exists(os.path.join(L4D2_ADDONS, m['file']))]
if _stale:
    print('[WARN] addons_scan.json 已过期：%d 个已扫描的 VPK 在磁盘上已不存在' % len(_stale))
    print('       例：%s' % '、'.join(_stale[:3]))
    print('       → 请先跑「刷新表格.bat」（标准流程第 9 步）再重跑本推演。')
    print()
lmap = json.load(open(os.path.join(W, 'label_map.json'), encoding='utf-8'))

# rv_paths.json 是按「文件名」索引的：改名后必须重跑，否则查不到路径，
# 就会漏掉 maps/*.bsp 与 .mdl 的判定（实测会把地图包误判成物件域）。
# 这里自愈：缺失、或比 addons_scan.json 旧，就自动重跑 rv_dump_paths.mjs。
_RVP = os.path.join(W, 'rv_paths.json')
_SCAN = os.path.join(W, 'addons_scan.json')
if (not os.path.exists(_RVP)) or (os.path.exists(_SCAN) and os.path.getmtime(_RVP) < os.path.getmtime(_SCAN)):
    import subprocess
    _node = os.environ.get('L4D2_NODE') or 'node'
    try:
        with open(os.path.join(os.path.dirname(W), 'config.json'), encoding='utf-8') as _f:
            _node = json.load(_f).get('nodePath') or _node
    except Exception:
        pass
    _dj = os.path.join(W, 'rv_dump_paths.mjs')
    if os.path.exists(_dj):
        print('[info] rv_paths.json 缺失或过期 → 自动重跑 rv_dump_paths.mjs')
        subprocess.run([_node, _dj], cwd=W, check=False)
paths = json.load(open(_RVP, encoding='utf-8'))

DOMAINS = ['生还者', '特殊感染者', '普通感染者', '主武器', '副武器', '近战',
           '投掷物', '消耗品', '场景道具', '固定武器·场景']
PREFIX = {'生还者': '生还者', '特殊感染者': '特感', '普通感染者': '普感', '主武器': '主武器',
          '副武器': '副武器', '近战': '近战', '投掷物': '投掷物', '消耗品': '消耗品',
          '场景道具': '场景道具', '固定武器·场景': '场景'}
WEAPON_FAMILY = ['主武器', '副武器', '近战', '投掷物', '消耗品', '场景道具']
EN_DOMAIN = {'生还者', '特殊感染者'}
CROSS = {
    '生还者骨架动画': '生还者', '生还者手势': '生还者', '生还者通用音效': '音效',
    '生还者语音': '音效', 'HUD 头像': '界面',
    'L4D1 版感染者模型': '特感', '感染者第一人称爪（第三人称视角模式）': '特感',
    '感染者手臂': '特感', '爆炸残肢': '特感',
    '女性感染者残肢模型': '普感', '普通感染者残肢模型': '普感', '血块与碎尸': '普感',
}
SUPER = {
    ('副武器', frozenset(['手枪（P220）', '格洛克手枪', '双持手枪'])): '手枪',
    ('主武器', frozenset(['泵动式霰弹枪', '镀铬霰弹枪'])): '霰弹枪',
    ('生还者', frozenset(['弗朗西斯', '弗朗西斯（轻量版）'])): 'Francis',
    ('特殊感染者', frozenset(['女巫', '新娘女巫'])): 'Witch',
    ('特殊感染者', frozenset(['呕吐者', '女呕吐者'])): 'Boomer',
    ('特殊感染者', frozenset(['坦克（巨兽）', '坦克（Cold Stream 版）'])): 'Tank',
    ('生还者', frozenset(['生还者骨架动画', '生还者手势'])): '静态表情',
}
CROSS_SHORT = {'生还者骨架动画': '骨架动画', '生还者手势': '手势', '生还者通用音效': '通用音效',
               '生还者语音': '语音', 'HUD 头像': 'HUD头像', '感染者第一人称爪（第三人称视角模式）': '第一人称爪',
               '感染者手臂': '手臂', '女性感染者残肢模型': '女性残肢', '普通感染者残肢模型': '残肢模型'}
MODEL_EXT = re.compile(r'\.(mdl|vvd|vtx|phy|ani)$', re.I)
BIG = 500

wb = load_workbook(XLSX, data_only=True)
ROWINFO, ALLROWS = {}, []
for s in DOMAINS:
    ws = wb[s]
    hdr = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    ci = {h: i + 1 for i, h in enumerate(hdr) if h}
    for r in range(3, ws.max_row + 1):
        cn = ws.cell(row=r, column=ci['中文名']).value
        en = ws.cell(row=r, column=ci['英文名']).value if '英文名' in ci else None
        ROWINFO[(s, cn)] = {'中文名': cn, '英文名': en}
        ALLROWS.append((s, cn, str(en or '')))

def sht(lab):
    v = lmap.get(lab); return v[0].split('::')[0] if v else None
def row_of(lab):
    v = lmap.get(lab); return v[0].split('::')[1] if v else None

def sound_domain(lab):
    if not lab.startswith('音效·'):
        return None
    tail = lab[3:]
    c = [(s, cn) for s, cn, _ in ALLROWS if s in WEAPON_FAMILY
         and (cn.startswith(tail) or tail.startswith(cn))]
    c.sort(key=lambda x: len(x[1]))
    return c[0][0] if c else None

plans, review = [], []
for m in scan:
    f = m['file']
    if f.startswith('cfhd'):
        continue
    stem = os.path.splitext(f)[0]
    tg = m.get('targets') or {}
    P = paths.get(f, [])
    total = sum(tg.values()) or 1
    has_model = any(MODEL_EXT.search(p) for p in P)
    oldpre = stem.split('-', 1)[0] if '-' in stem else None
    unk = (tg.get('未归类') or 0) / total

    # S1
    if any(re.search(r'\.bsp$|\.nav$', p, re.I) for p in P):
        plans.append(dict(stem=stem, pre='地图', tgt='', why='S1 含 bsp/nav', oldpre=oldpre)); continue
    # S5a 未归类过半
    if unk > 0.5:
        plans.append(dict(stem=stem, pre=oldpre or '其它', tgt='', why='S5 未归类 %.0f%%' % (unk*100), oldpre=oldpre))
        review.append((stem, 'S5 未归类占比 %.0f%%，保留原前缀' % (unk*100), '')); continue

    # 物件标签
    obj = {}
    for lab, n in tg.items():
        s = sht(lab)
        if s in DOMAINS:
            obj.setdefault(s, set()).add(row_of(lab))
    # 邻接标签（只在该 mod 含模型文件时参与）
    adj = {}
    if has_model:
        for lab, n in tg.items():
            if sht(lab) in DOMAINS:
                continue
            d = None
            if lab in ('材质·生还者贴图', '音效·生还者技能音效'):
                d = '生还者'
            elif lab in ('动画·感染者骨骼', '材质·感染者贴图'):
                d = '特殊感染者'
            elif lab.startswith('语音·'):
                d = '生还者'
            elif lab == '材质·武器贴图':
                fam = [k for k in obj if k in WEAPON_FAMILY]
                d = fam[0] if len(fam) == 1 else None
            else:
                d = sound_domain(lab)
            if d:
                adj[d] = adj.get(d, 0) + n

    def domct(d):
        n = 0
        for lab, k in tg.items():
            if sht(lab) == d:
                n += k
        return n + adj.get(d, 0)

    stage = None
    if obj:
        cands = {d: domct(d) for d in obj}
        d0 = min(cands, key=lambda d: DOMAINS.index(d))
        # 取域优先级最高者
        best_d = min(obj, key=lambda d: DOMAINS.index(d))
        ratio = domct(best_d) / total
        if not has_model:
            allr = set().union(*obj.values())
            if allr <= set(CROSS):
                r0 = max(allr, key=lambda r: tg.get(next((l for l in tg if row_of(l) == r), ''), 0))
                pre0 = CROSS[r0]
                plans.append(dict(stem=stem, pre=pre0, tgt='', why='S3b 纯横切行(无模型) → %s-' % pre0, oldpre=oldpre))
                review.append((stem, '纯横切行且无模型，按「%s」取 %s-' % (r0, pre0), ''))
                continue
            stage = 'S3'
        elif total >= BIG and ratio < 0.05:
            plans.append(dict(stem=stem, pre=oldpre or '其它', tgt='', why='S5 大而散(总%d条,域%.1f%%)' % (total, ratio*100), oldpre=oldpre))
            review.append((stem, 'S5 大而散：总 %d 条路径、域占比仅 %.1f%%，保留原前缀' % (total, ratio*100), ''))
            continue
        else:
            stage = 'S2'

    if stage:
        rows_in = sorted(obj[best_d])
        main = [r for r in rows_in if r not in CROSS]
        pool = main or rows_in
        pick = max(pool, key=lambda r: tg.get(next((l for l in tg if row_of(l) == r), ''), 0))
        key = (best_d, frozenset(pool))
        if key in SUPER:
            tgt = SUPER[key]
        else:
            tgt = pick
            if not main:
                tgt = CROSS_SHORT.get(pick, pick)
                review.append((stem, '只命中横切行，中段取「%s」' % tgt, '、'.join(rows_in)))
            elif len(pool) > 1:
                review.append((stem, '同域多行，中段取占比最高的「%s」' % pick, '、'.join(pool)))
        info = ROWINFO.get((best_d, pick)) or {}
        if best_d in EN_DOMAIN and info.get('英文名') and not str(info['英文名']).startswith('—'):
            tgt = info['英文名']
        plans.append(dict(stem=stem, pre=PREFIX[best_d], tgt=tgt,
                          why='%s 物件域=%s(%.0f%%)' % (stage, best_d, domct(best_d)/total*100), oldpre=oldpre))
        if stage == 'S3':
            review.append((stem, 'S3 中置信（无模型文件，域占比 %.0f%%）' % (domct(best_d)/total*100), '物件域=%s' % best_d))
        continue

    # S4
    typ = {}
    for lab, nn in tg.items():
        if sht(lab) in DOMAINS:
            continue
        typ[lab.split('·')[0]] = typ.get(lab.split('·')[0], 0) + nn
    TYPEPRE = {'材质': '材质', '界面': '界面', '特效': '特效', '脚本': '脚本',
               '音效': '音效', '语音': '音效', '地图': '地图', '模型': '模型',
               '手电筒': '手电筒', '管理员插件': '管理员插件'}
    got = None
    # 2026-10-03 新增的两类需要**占比门槛**：防止「顺带带了一张手电筒贴图 / 一个 admin 脚本」
    # 的 mod 被抢走身份。实测 材质-ESC菜单 Kokomi自用版 4 条路径里 2 条是手电筒贴图（50%），
    # 但它其实是菜单 mod —— 故收窄规则之余再加门槛。真实成员占比：手电筒 100%、管理员插件 79%。
    NEWCAT_MIN = 0.25
    # 管理插件优先：它靠**入口文件**定身份（多数路径仍是 scripts/vscripts/ 通用库），
    # 所以门槛只需挡住「顺带带一个 admin 模块」的脚本合集包。
    # 两道判定都放在 S4 内、优先于模型/素材，避免误伤地图（地图走 S1，不会到这儿）。
    if typ.get('管理员插件', 0) / total >= NEWCAT_MIN:
        got = '管理员插件'
    elif typ.get('手电筒', 0) / total >= NEWCAT_MIN:
        got = '手电筒'
    elif has_model and (typ.get('模型', 0) / total) >= 0.20:
        got = '模型'
    else:
        for pre, _ in sorted(typ.items(), key=lambda kv: -kv[1]):
            if pre in TYPEPRE and pre != '模型':
                got = TYPEPRE[pre]; break
        if got is None and typ.get('模型'):
            got = '模型'
    if got is None:
        plans.append(dict(stem=stem, pre=oldpre or '其它', tgt='', why='S5 无可用标签', oldpre=oldpre))
        review.append((stem, 'S5 无可用标签，保留原前缀', ''))
        continue
    plans.append(dict(stem=stem, pre=got, tgt='', why='S4 资源类型', oldpre=oldpre))

# ---- 多件套归并：`A` 与 `A-后缀` 共用主体前缀 ----
byname = {p['stem']: p for p in plans}
for p in list(plans):
    for main in byname:
        if p['stem'] != main and p['stem'].startswith(main + '-'):
            byname[main]['members'] = byname[main].get('members', []) + [p['stem']]

def newname(p):
    """<旧前缀>-<旧替换对象>-<标题>  ->  <新前缀>-<新替换对象>-<标题>
    幂等保证：若当前名已经以「新前缀-新替换对象」开头，原样保留（不再动）。
    这样即使新替换对象自带短横（AK-47 / M-16），重复跑也不会被切坏。"""
    stem, pre, tgt = p['stem'], p['pre'], p['tgt']
    rest = stem.split('-', 1)[1] if '-' in stem else stem
    want = ('%s-%s' % (pre, tgt)) if tgt else pre
    if stem == want or stem.startswith(want + '-'):
        return stem
    if not tgt:
        return '%s-%s' % (pre, rest)
    # 旧名里「旧替换对象」是旧前缀之后的第一个短横段；旧替换对象本身不含短横
    body = rest.split('-', 1)[1] if '-' in rest else rest
    return '%s-%s-%s' % (pre, tgt, body) if body else '%s-%s' % (pre, tgt)

for p in plans:
    p['new'] = newname(p)
# 成员跟随主体
for p in plans:
    for main, mp in byname.items():
        if p['stem'] != main and p['stem'].startswith(main + '-'):
            suf = p['stem'][len(main):]
            p['new'] = mp['new'] + suf
            p['why'] += '（多件套跟随主体）'

# ---- 规则例外：用户拍板的取舍，重跑时必须保留，不得按规则输出覆盖 ----
_ex = os.path.join(W, 'name_exempt.json')
if os.path.exists(_ex):
    _EX = json.load(open(_ex, encoding='utf-8'))
    for p in plans:
        for e in _EX:
            if p['stem'] in (e.get('old'), e.get('name')):
                if p['new'] != e['name']:
                    review.append((p['stem'], '命中规则例外 → 固定为 %s' % e['name'], e.get('reason', '')[:50]))
                p['new'] = e['name']
                p['why'] += '（规则例外）'

order = ['生还者', '特感', '普感', '主武器', '副武器', '近战', '投掷物', '消耗品',
         '场景道具', '场景', '模型', '地图', '材质', '界面', '特效', '脚本', '音效',
         '手电筒', '管理员插件', '其它']
plans.sort(key=lambda p: (order.index(p['pre']) if p['pre'] in order else 99, p['stem']))

lines = ['# MOD 重命名干跑计划（工作流 v1.0）\n',
         '> S1 地图 → S2 物件域(有物件标签+有模型) → S3 物件域(无模型=贴图/音效类) → S4 资源类型 → S5 保留原前缀\n',
         '| # | 现名 | 新名 | 依据 |', '|---|---|---|---|']
cnt, chg = {}, 0
for i, p in enumerate(plans, 1):
    cnt[p['pre']] = cnt.get(p['pre'], 0) + 1
    if p['new'] != p['stem']:
        chg += 1
    lines.append('| %d | `%s` | `%s` | %s |' % (i, p['stem'], p['new'], p['why']))
lines.append('\n## 待复核（%d 条）\n' % len(review))
lines.append('| mod | 原因 | 细节 |\n|---|---|---|')
for s, r, d in review:
    lines.append('| `%s` | %s | %s |' % (s, r, d))
open(OUT, 'w', encoding='utf-8').write('\n'.join(lines))

print('共 %d 个，需改名 %d 个' % (len(plans), chg))
for k in order:
    if cnt.get(k):
        print('   %-8s %d' % (k + '-', cnt[k]))
print()
print('待复核 %d 条：' % len(review))
for s, r, d in review:
    print('   %-58s %s %s' % (s, r, d))
